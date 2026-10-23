#!/usr/bin/env python3
"""Embed the Assessment 2 diagrams into the report under docs/."""

from __future__ import annotations

import os
import struct
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/assessment2/MSE907-Assessment2-WIP-zh.docx"
FIGURES = ROOT / "docs/assessment2/figures"

FIGURE_SPECS = (
    {
        "anchor": "8. 使用的工具与技术/Tools and Technologies Used",
        "caption": "图 1 数据流与证据边界 / Figure 1. Data flow and evidence boundary",
        "image": "figure-1-data-flow.png",
    },
    {
        "anchor": "9.2 不可变版本、契约和可验证性",
        "caption": "图 2 高层架构与信任边界 / Figure 2. High-level architecture and trust boundaries",
        "image": "figure-2-architecture.png",
    },
    {
        "anchor": "10. 实施计划与进展/Implementation Plan and Progress",
        "caption": "图 3 UID 保护清理状态机 / Figure 3. UID-guarded cleanup state machine",
        "image": "figure-3-cleanup-state-machine.png",
    },
    {
        "anchor": "11. 评价或测试策略/Evaluation or Testing Strategy (if applicable)",
        "caption": "图 4 含失败分支的预览生命周期 / Figure 4. Preview lifecycle with failure branches",
        "image": "figure-4-lifecycle-workflow.png",
    },
)


def find_paragraph(document: Document, anchor: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == anchor:
            return paragraph
    raise ValueError(f"could not find report heading: {anchor}")


def png_dimensions(path: Path) -> tuple[int, int]:
    """Read a PNG's pixel dimensions without adding an image dependency."""
    header = path.read_bytes()[:24]
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"not a PNG file: {path}")
    return struct.unpack(">II", header[16:24])


def add_figure(document: Document, spec: dict[str, str]) -> None:
    caption = spec["caption"]
    if any(paragraph.text.strip() == caption for paragraph in document.paragraphs):
        return

    anchor = find_paragraph(document, spec["anchor"])
    image_path = FIGURES / spec["image"]
    if not image_path.is_file():
        raise FileNotFoundError(image_path)

    image_paragraph = anchor.insert_paragraph_before(style="Normal")
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_together = True
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.0))
    caption_paragraph = anchor.insert_paragraph_before(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.paragraph_format.keep_with_next = False
    caption_paragraph.paragraph_format.space_after = Pt(8)


def update_figure_sizes(document: Document) -> None:
    """Keep each existing diagram frame proportional to its rendered PNG."""
    if len(document.inline_shapes) < len(FIGURE_SPECS):
        raise ValueError("the report does not contain all four diagram frames")

    width = Inches(6.0)
    for shape, spec in zip(document.inline_shapes, FIGURE_SPECS):
        image_width, image_height = png_dimensions(FIGURES / spec["image"])
        shape.width = width
        shape.height = int(width * image_height / image_width)


def update_figure_paragraphs(document: Document) -> None:
    """Keep each diagram and its caption together across page breaks."""
    captions = {spec["caption"] for spec in FIGURE_SPECS}
    for paragraph in document.paragraphs:
        if paragraph._p.xpath(".//wp:inline"):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.strip() in captions:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False


def refresh_embedded_images(path: Path) -> None:
    """Replace diagram media without changing the report's paragraph layout."""
    replacements = {
        f"word/media/image{index}.png": (FIGURES / spec["image"]).read_bytes()
        for index, spec in enumerate(FIGURE_SPECS, start=1)
    }
    with tempfile.NamedTemporaryFile(
        dir=path.parent, prefix=f".{path.stem}-", suffix=".tmp.docx", delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)

    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination:
            for member in source.infolist():
                data = replacements.get(member.filename)
                if data is None:
                    data = source.read(member.filename)
                destination.writestr(member, data)
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def main() -> None:
    document = Document(SOURCE)
    for spec in FIGURE_SPECS:
        add_figure(document, spec)
    update_figure_sizes(document)
    update_figure_paragraphs(document)
    document.save(SOURCE)
    refresh_embedded_images(SOURCE)
    print(f"updated {SOURCE}")


if __name__ == "__main__":
    main()
